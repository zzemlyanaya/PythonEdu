import re

from dataclasses import dataclass
from docx import Document
from docx.shared import Pt


@dataclass
class EncrypterArgs:
    doc_name: str
    x: float


class HideousEncrypter:
    def __init__(self, args: EncrypterArgs):
        self.x = args.x

        self.original_file = './original/' + args.doc_name
        self.encoded_file = './encoded/encoded_' + args.doc_name
        self.decode_file = './encoded/decode_file.docx'
        self.cypher_file = './original/cipher.txt'

    def work(self):
        with open(self.cypher_file, "r") as file:
            cipher = file.read().lower()

        doc = Document(self.original_file)
        paragraphs = doc.paragraphs
        text = "\n".join([p.text for p in paragraphs])

        if len(cipher) > len(text):
            raise RuntimeError("Secret message is too long!")

        label = -1
        markers = []
        for c in cipher:
            upper = text.find(c.upper(), label + 1)
            lower = text.find(c, label + 1)

            if (upper == -1) and (lower == -1):
                markers.append(-1)
            elif (upper == -1) and (lower != -1):
                markers.append(lower)
            elif (upper != -1) and (lower == -1):
                markers.append(upper)
            else:
                final = min(lower, upper)
                markers.append(final)
                label = final

        if -1 in markers:
            raise RuntimeError("Can't hide the message!")
        else:
            markers_len = len(markers)
            encoded = Document()
            paragraphs = encoded.add_paragraph(text[0:markers[0]])
            for i in range(markers_len):
                run = paragraphs.add_run(text[markers[i]])
                run.font.size = Pt(self.x * 11)
                if i == markers_len - 1:
                    paragraphs.add_run(text[markers[i]+1:])
                else:
                    paragraphs.add_run(text[markers[i]+1:markers[i+1]])
            encoded.save(self.encoded_file)

    def decode(self):
        doc = Document(self.decode_file)
        paragraphs = doc.paragraphs
        cipher = ''
        for p in paragraphs:
            for run in p.runs:
                if run.font.size == Pt(11 * self.x):
                    cipher += run.text
        print(cipher)
